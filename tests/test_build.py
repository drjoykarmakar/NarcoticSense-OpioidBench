import csv
import json
from pathlib import Path

from opioidbench.cli import main

ROOT = Path(__file__).parents[1]


def test_data_build_writes_versioned_audit_and_inventory(tmp_path: Path) -> None:
    output = tmp_path / "dataset"
    code = main(
        [
            "data",
            "build",
            "--input",
            str(ROOT / "tests" / "fixtures" / "synthetic_swgdrug_layout.hpj"),
            "--source",
            "synthetic_fixture",
            "--output",
            str(output),
            "--dataset-version",
            "test-v1",
            "--taxonomy",
            str(ROOT / "metadata" / "opioid_taxonomy.csv"),
            "--registry",
            str(ROOT / "metadata" / "source_registry.yaml"),
        ]
    )
    assert code == 0
    manifest = json.loads((output / "dataset_manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((output / "dataset_audit.json").read_text(encoding="utf-8"))
    assert manifest["dataset_version"] == "test-v1"
    assert manifest["number_spectra"] == 2
    assert audit["number_unique_compounds"] == 2
    assert audit["placeholder_cas_records"] == 2
    assert audit["taxonomy_complete"] is False
    with (output / "compounds.csv").open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2
    assert {row["preferred_name"] for row in rows} == {
        "Synthetic Compound Alpha",
        "Synthetic Compound Beta",
    }


def test_data_build_records_companion_metadata(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    companion = tmp_path / "compound_list.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1).value = "Compound Name (Red Indicates new entry in v 3.14)"
    sheet.cell(2, 1).value = "Synthetic Compound Alpha"
    sheet.cell(3, 1).value = "Synthetic Compound Beta"
    sheet.cell(3, 1).font = Font(color="FF0000")
    workbook.save(companion)

    output = tmp_path / "dataset_with_companion"
    code = main(
        [
            "data",
            "build",
            "--input",
            str(ROOT / "tests" / "fixtures" / "synthetic_swgdrug_layout.hpj"),
            "--source",
            "synthetic_fixture",
            "--output",
            str(output),
            "--dataset-version",
            "test-v2",
            "--taxonomy",
            str(ROOT / "metadata" / "opioid_taxonomy.csv"),
            "--compound-list",
            str(companion),
            "--registry",
            str(ROOT / "metadata" / "source_registry.yaml"),
        ]
    )
    assert code == 0
    manifest = json.loads((output / "dataset_manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((output / "dataset_audit.json").read_text(encoding="utf-8"))
    assert manifest["companion_file"] == "compound_list.xlsx"
    assert manifest["companion_file_sha256"]
    assert audit["companion_name_relation_counts"] == {"exact": 2}
    assert audit["new_in_source_version_counts"] == {"3.14": 1}
    review = (output / "companion_name_review.csv").read_text(encoding="utf-8")
    assert review.splitlines() == [
        "spectrum_id,source_record_id,compound_name_hpj,compound_name_companion,"
        "name_relation,cas,new_in_source_version"
    ]

    records = [
        json.loads(line)
        for line in (output / "normalized_records.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert records[0]["compound_name_companion"] == "Synthetic Compound Alpha"
    assert records[0]["compound_name_resolution"] == "exact"
    assert records[1]["new_in_source_version"] == "3.14"
