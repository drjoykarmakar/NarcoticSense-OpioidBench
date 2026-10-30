from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from opioidbench.data.companion import (
    align_swgdrug_compound_list,
    classify_name_relation,
    preferred_source_name,
    read_swgdrug_compound_list,
)
from opioidbench.data.schema import RawSpectrum, SourceProvenance, SpectrumPeak


def _raw(name: str, record_id: str) -> RawSpectrum:
    return RawSpectrum(
        spectrum_id=f"synthetic:{record_id}",
        compound_name_raw=name,
        peaks=(SpectrumPeak(50.0, 100.0),),
        provenance=SourceProvenance(
            source="swgdrug",
            source_version="3.14",
            source_record_id=record_id,
            retrieval_date=date(2026, 9, 29),
            source_url=None,
            license=None,
        ),
    )


def _write_companion(path: Path, names: list[str], *, red_index: int | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SWGDRUG formula exact mass"
    sheet.cell(1, 1).value = "Compound Name (Red Indicates new entry in v 3.14)"
    for index, name in enumerate(names, start=2):
        cell = sheet.cell(index, 1)
        cell.value = name
        if red_index is not None and index - 2 == red_index:
            cell.font = Font(color="FF0000")
    workbook.save(path)


def test_companion_alignment_preserves_source_disagreement(tmp_path: Path) -> None:
    long_name = "N-" + "A" * 70
    truncated = long_name[:65]
    hpj_longer = 'Reduced Pyrilamine by "Nazi" method'
    companion_shorter = "Reduced Pyrilamine by"
    names = ["Exact Name", long_name, companion_shorter]
    workbook = tmp_path / "compound_list.xlsx"
    _write_companion(workbook, names, red_index=1)

    entries, version = read_swgdrug_compound_list(workbook)
    aligned, report = align_swgdrug_compound_list(
        [
            _raw("Exact Name", "1"),
            _raw(truncated, "2"),
            _raw(hpj_longer, "3"),
        ],
        entries,
    )

    assert version == "3.14"
    assert report.relation_counts == {
        "companion_shorter": 1,
        "exact": 1,
        "hpj_truncated": 1,
    }
    assert report.number_new_in_source_version == 1
    assert preferred_source_name(aligned[1]) == long_name
    assert preferred_source_name(aligned[2]) == hpj_longer
    assert aligned[2].metadata["COMPOUND_LIST_NAME"] == companion_shorter


def test_companion_alignment_rejects_unrelated_names(tmp_path: Path) -> None:
    workbook = tmp_path / "compound_list.xlsx"
    _write_companion(workbook, ["Completely Different"])
    entries, _ = read_swgdrug_compound_list(workbook)
    with pytest.raises(ValueError, match="Unsafe companion alignment"):
        align_swgdrug_compound_list([_raw("Original Name", "1")], entries)


def test_name_relation_requires_observed_65_character_truncation_pattern() -> None:
    assert classify_name_relation("Short", "Shorter") == "mismatch"
    hpj = "A" * 65
    assert classify_name_relation(hpj, hpj + " tail") == "hpj_truncated"


def _write_exact_mass_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SWGDRUG formula exact mass"
    sheet.cell(1, 1).value = "Monoisotopic masses"
    sheet.cell(2, 1).value = (
        "All masses are calculated from the molecular formula present in the SWGDRUG "
        "GCMS library using AGILENT PCDL software and are unveriified."
    )
    sheet.cell(7, 1).value = "SWGDrug Library version:"
    sheet.cell(7, 2).value = "SWGDRUG MS Library Version 3.14"
    for column, value in enumerate(["Formula", "Compound Name", "Mass", "M+proton"], start=1):
        sheet.cell(8, column).value = value
    sheet.cell(9, 1).value = "C2H6O"
    sheet.cell(9, 2).value = "Exact Name"
    sheet.cell(9, 3).value = 46.041864812
    sheet.cell(9, 4).value = 47.049164812
    sheet.cell(10, 1).value = "C3H8O"
    sheet.cell(10, 2).value = "Another Name"
    sheet.cell(10, 3).value = 60.057514876
    workbook.save(path)


def test_exact_mass_companion_alignment(tmp_path: Path) -> None:
    from opioidbench.data.companion import (
        align_swgdrug_exact_mass_list,
        read_swgdrug_exact_mass_list,
    )

    workbook = tmp_path / "exact_mass.xlsx"
    _write_exact_mass_workbook(workbook)
    entries, version = read_swgdrug_exact_mass_list(workbook)
    spectra = [
        _raw("Exact Name", "1"),
        _raw("Another Name", "2"),
    ]
    spectra = [
        RawSpectrum(
            spectrum_id=item.spectrum_id,
            compound_name_raw=item.compound_name_raw,
            peaks=item.peaks,
            provenance=item.provenance,
            metadata={"MOLFORM": formula},
        )
        for item, formula in zip(spectra, ["C2H6O", "C3H8O"], strict=True)
    ]

    aligned, report = align_swgdrug_exact_mass_list(spectra, entries)

    assert version == "3.14"
    assert report.number_entries == 2
    assert report.number_formula_matches == 2
    assert report.number_m_plus_proton_values == 1
    assert report.number_m_plus_proton_missing == 1
    assert aligned[0].metadata["EXACT_MASS_LIST_MONOISOTOPIC_MASS"] == repr(46.041864812)
    assert "EXACT_MASS_LIST_M_PLUS_PROTON" not in aligned[1].metadata


def test_exact_mass_companion_rejects_formula_disagreement(tmp_path: Path) -> None:
    from opioidbench.data.companion import (
        align_swgdrug_exact_mass_list,
        read_swgdrug_exact_mass_list,
    )

    workbook = tmp_path / "exact_mass.xlsx"
    _write_exact_mass_workbook(workbook)
    entries, _ = read_swgdrug_exact_mass_list(workbook)
    spectrum = _raw("Exact Name", "1")
    spectrum = RawSpectrum(
        spectrum_id=spectrum.spectrum_id,
        compound_name_raw=spectrum.compound_name_raw,
        peaks=spectrum.peaks,
        provenance=spectrum.provenance,
        metadata={"MOLFORM": "C2H4O"},
    )
    with pytest.raises(ValueError, match="Formula disagreement"):
        align_swgdrug_exact_mass_list([spectrum], entries[:1])


def test_agilent_header_restores_full_source_comment(tmp_path: Path) -> None:
    import zipfile

    from opioidbench.data.companion import (
        align_swgdrug_agilent_header,
        read_swgdrug_agilent_header_archive,
    )

    archive_path = tmp_path / "agilent.zip"
    header = b"\x00".join(
        [
            b"SWGDRUG 3.",
            b"Exact Name",
            b"C2H6O",
            b"Source ABC with full provenance",
            b"SWGDRUG 3.",
            b"BDB",
            b"C11H15NO2",
            b"ENFSI test source",
        ]
    )
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("SWGDRUG 3.14.L/HEADER", header)

    first = _raw("Exact Name", "1")
    first = RawSpectrum(
        spectrum_id=first.spectrum_id,
        compound_name_raw=first.compound_name_raw,
        peaks=first.peaks,
        provenance=first.provenance,
        metadata={"MOLFORM": "C2H6O", "SAMPLE_DESCRIPTION": "Source ABC"},
    )
    second = _raw("BDB", "2")
    second = RawSpectrum(
        spectrum_id=second.spectrum_id,
        compound_name_raw=second.compound_name_raw,
        peaks=second.peaks,
        provenance=second.provenance,
        metadata={"MOLFORM": "C11H15NO2", "SAMPLE_DESCRIPTION": "ENFSI test source"},
    )

    entries, member, digest = read_swgdrug_agilent_header_archive(archive_path)
    aligned, report = align_swgdrug_agilent_header(
        [first, second],
        entries,
        archive_member=member,
        header_sha256=digest,
    )

    assert report.number_entries == 2
    assert report.number_extended_source_comments == 1
    assert report.number_short_names_omitted_from_printable_scan == 1
    assert report.source_comment_relation_counts == {"exact": 1, "hpj_truncated": 1}
    assert aligned[0].metadata["SOURCE_COMMENT_FULL"] == "Source ABC with full provenance"
    assert aligned[1].metadata["SOURCE_COMMENT_FULL"] == "ENFSI test source"


def test_agilent_header_rejects_comment_disagreement(tmp_path: Path) -> None:
    import zipfile

    from opioidbench.data.companion import (
        align_swgdrug_agilent_header,
        read_swgdrug_agilent_header_archive,
    )

    archive_path = tmp_path / "agilent.zip"
    header = b"\x00".join(
        [b"SWGDRUG 3.", b"Exact Name", b"C2H6O", b"Different provenance"]
    )
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("SWGDRUG 3.14.L/HEADER", header)
    spectrum = _raw("Exact Name", "1")
    spectrum = RawSpectrum(
        spectrum_id=spectrum.spectrum_id,
        compound_name_raw=spectrum.compound_name_raw,
        peaks=spectrum.peaks,
        provenance=spectrum.provenance,
        metadata={"MOLFORM": "C2H6O", "SAMPLE_DESCRIPTION": "Source ABC"},
    )
    entries, member, digest = read_swgdrug_agilent_header_archive(archive_path)
    with pytest.raises(ValueError, match="source-comment alignment failed"):
        align_swgdrug_agilent_header(
            [spectrum], entries, archive_member=member, header_sha256=digest
        )
